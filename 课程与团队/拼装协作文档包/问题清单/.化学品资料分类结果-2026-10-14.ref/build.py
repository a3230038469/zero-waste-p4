# -*- coding: utf-8 -*-
"""化学品资料分类结果 → 单表 xlsx（记录型清单：冻结表头 + 筛选 + 轻量灰）"""
try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEAD_BG = xl_color("#F2F2F2")     # 轻量灰：表头底
XL_HEAD_FG = xl_color("#1A1A1A")     # 表头字
XL_BORDER = xl_color("#BFBFBF")

THIN = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SRC = r"D:/projects/_chem_result.tsv"
OUT = r"C:/Users/a3230/Desktop/问题清单/化学品资料分类结果-2026-10-14.xlsx"

# 读数据
table = []
with open(SRC, encoding="utf-8") as f:
    for line in f:
        line = line.rstrip("\n")
        if not line:
            continue
        table.append(line.split("\t"))

header, data = table[0], table[1:]
assert header == ["序号", "资料名", "发布机构", "发布年份", "类型", "主题"], header
assert len(data) == 123, len(data)

wb = Workbook()
ws = wb.active
ws.title = "化学品资料分类"
wb.properties.title = "化学品资料分类结果"

# 表头（第 1 行）
for c, name in enumerate(header, start=1):
    cell = ws.cell(row=1, column=c, value=name)
    cell.font = Font(bold=True, color=XL_HEAD_FG, size=11)
    cell.fill = PatternFill("solid", fgColor=XL_HEAD_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = BORDER

# 数据（第 2 行起）
for i, row in enumerate(data):
    r = 2 + i
    seq, title, org, year, typ, topic = row
    vals = [int(seq), title, org, (int(year) if year.strip().isdigit() else year), typ, topic]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.border = BORDER
        if c == 1:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 4:
            cell.number_format = "0"
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif c == 6:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=(c == 2))

# 列宽：序号窄、资料名宽、其余按内容
widths = {"A": 6, "B": 54, "C": 30, "D": 10, "E": 12, "F": 24}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

last_row = 1 + len(data)
ws.freeze_panes = "A2"                      # 冻结表头
ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{last_row}"   # 开启筛选
ws.sheet_view.zoomScale = 100

wb.save(OUT)
print("已生成:", OUT)
print("数据行:", len(data), "| 列:", len(header), "| 冻结: A2 | 筛选:", ws.auto_filter.ref)
